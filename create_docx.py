import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document():
    md_file_path = r"C:\Users\Ykmacho\.gemini\antigravity-ide\brain\8281573a-b7df-4985-b812-e0bf73fa70ba\plan_implementacion_gss.md"
    docx_file_path = r"C:\GSS\Plan_de_Implementacion_GSS.docx"

    document = Document()
    
    style = document.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    with open(md_file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    
    for line in lines:
        line = line.strip()
        if not line:
            document.add_paragraph()
            continue
            
        if line == '---' or line == '&nbsp;':
            continue

        if line.startswith('# '):
            h = document.add_heading(line[2:], level=1)
            h.style.font.name = 'Times New Roman'
        elif line.startswith('## '):
            h = document.add_heading(line[3:], level=1)
            h.style.font.name = 'Times New Roman'
        elif line.startswith('### '):
            h = document.add_heading(line[4:], level=2)
            h.style.font.name = 'Times New Roman'
        elif line.startswith('- '):
            p = document.add_paragraph(style='List Bullet')
            text = line[2:]
            parts = re.split(r'\*\*(.*?)\*\*', text)
            for i, part in enumerate(parts):
                if i % 2 == 1:
                    p.add_run(part).bold = True
                else:
                    p.add_run(part)
        elif line.startswith('|'):
            p = document.add_paragraph(line)
            p.style.font.name = 'Courier New'
            p.style.font.size = Pt(10)
        else:
            p = document.add_paragraph()
            parts = re.split(r'\*\*(.*?)\*\*', line)
            for i, part in enumerate(parts):
                if i % 2 == 1:
                    p.add_run(part).bold = True
                else:
                    p.add_run(part)

    document.save(docx_file_path)
    print(f"Document saved to {docx_file_path}")

if __name__ == '__main__':
    create_document()
