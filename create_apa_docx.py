import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_document():
    md_file_path = r"C:\Users\Ykmacho\.gemini\antigravity-ide\brain\8281573a-b7df-4985-b812-e0bf73fa70ba\plan_implementacion_gss.md"
    docx_file_path = r"C:\GSS\Plan_de_Implementacion_GSS_APA.docx"

    document = Document()

    sections = document.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = document.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 2.0
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)
    
    apa_style = document.styles.add_style('APA Paragraph', WD_STYLE_TYPE.PARAGRAPH)
    apa_style.base_style = style
    apa_style.paragraph_format.first_line_indent = Inches(0.5)

    h1_style = document.styles['Heading 1']
    h1_style.font.name = 'Times New Roman'
    h1_style.font.size = Pt(12)
    h1_style.font.bold = True
    h1_style.font.color.rgb = RGBColor(0, 0, 0)
    h1_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1_style.paragraph_format.line_spacing = 2.0
    h1_style.paragraph_format.space_before = Pt(0)
    h1_style.paragraph_format.space_after = Pt(0)
    h1_style.paragraph_format.first_line_indent = Inches(0)
    
    h2_style = document.styles['Heading 2']
    h2_style.font.name = 'Times New Roman'
    h2_style.font.size = Pt(12)
    h2_style.font.bold = True
    h2_style.font.color.rgb = RGBColor(0, 0, 0)
    h2_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2_style.paragraph_format.line_spacing = 2.0
    h2_style.paragraph_format.space_before = Pt(0)
    h2_style.paragraph_format.space_after = Pt(0)
    h2_style.paragraph_format.first_line_indent = Inches(0)

    with open(md_file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    
    in_title_page = False
    in_table = False
    table_data = []

    def format_text(p, text):
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|.*?)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                p.add_run(part[2:-2]).bold = True
            elif part.startswith('*') and part.endswith('*'):
                p.add_run(part[1:-1]).italic = True
            elif part.startswith('') and part.endswith(''):
                run = p.add_run(part[1:-1])
                run.font.name = 'Courier New'
            else:
                p.add_run(part)

    def set_cell_border(tc, **kwargs):
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right'):
            edge_data = kwargs.get(edge)
            if edge_data:
                tag = 'w:{}'.format(edge)
                element = OxmlElement(tag)
                for key in ['val', 'color', 'space', 'sz']:
                    if key in edge_data:
                        element.set(qn('w:{}'.format(key)), str(edge_data[key]))
                tcBorders.append(element)
        tcPr.append(tcBorders)

    def render_table(data):
        if not data: return
        clean_data = []
        for row in data:
            if '---' in row: continue
            cells = [c.strip() for c in row.strip('|').split('|') if c.strip() != '']
            if cells:
                clean_data.append(cells)
            
        if not clean_data: return
        
        try:
            table = document.add_table(rows=len(clean_data), cols=len(clean_data[0]))
            table.style = 'Table Grid'
            
            for i, row in enumerate(clean_data):
                for j, cell_text in enumerate(row):
                    if j >= len(table.columns): continue
                    cell = table.cell(i, j)
                    p = cell.paragraphs[0]
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(0)
                    format_text(p, cell_text)
                    
                    set_cell_border(
                        cell._tc,
                        top={'val': 'single', 'sz': '12'} if i == 0 else ({'val': 'single', 'sz': '6'} if i == 1 else {'val': 'nil'}),
                        bottom={'val': 'single', 'sz': '12'} if i == len(clean_data)-1 else {'val': 'nil'},
                        left={'val': 'nil'},
                        right={'val': 'nil'}
                    )
            document.add_paragraph()
        except Exception as e:
            print("Error parsing table", e)

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        
        if line.startswith('|'):
            if not in_table:
                in_table = True
                table_data = []
            table_data.append(line)
            continue
        else:
            if in_table:
                render_table(table_data)
                in_table = False
                table_data = []

        if not line:
            # Skip empty lines
            continue
            
        if line == '---':
            continue
        if line == '&nbsp;':
            document.add_paragraph()
            continue

        if line == '## Portada':
            in_title_page = True
            for _ in range(4):
                document.add_paragraph()
            continue
            
        if line == '## Tabla de Contenido':
            document.add_page_break()
            in_title_page = False
            
        if in_title_page:
            p = document.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            
            if line.startswith('**') and line.endswith('**'):
                line = line[2:-2]
                run = p.add_run(line)
                run.bold = True
            else:
                p.add_run(line)
            continue

        if line.startswith('# '):
            document.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            document.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            document.add_heading(line[4:], level=2)
        elif line.startswith('#### '):
            document.add_heading(line[5:], level=3)
        elif line.startswith('- '):
            p = document.add_paragraph(style='List Bullet')
            p.paragraph_format.line_spacing = 2.0
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Inches(0)
            format_text(p, line[2:])
        elif line.startswith('`') or line.startswith('+') or line.startswith('¦') or line.startswith('+'):
            p = document.add_paragraph(style='Normal')
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(line)
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
        else:
            p = document.add_paragraph(style='APA Paragraph')
            format_text(p, line)

    if in_table:
        render_table(table_data)

    document.save(docx_file_path)
    print("APA document saved successfully.")

if __name__ == '__main__':
    create_document()
